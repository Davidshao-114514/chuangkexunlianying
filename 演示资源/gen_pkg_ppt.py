# -*- coding: utf-8 -*-
"""开源学习包演示 PPT —— 磁吸焊接辅助手电"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN

BASE = r"C:\Users\35464\Desktop\创客训练营"
OUT = os.path.join(BASE, "演示资源")
os.makedirs(OUT, exist_ok=True)
FNAME = os.path.join(OUT, "磁吸焊接辅助手电_开源学习包_演示PPT.pptx")
BLOCK = os.path.join(BASE, "docs", "系统框图.png")

BLUE = C(31, 56, 100); DGRAY = C(55, 66, 88); LGRAY = C(128, 138, 155)
RED = C(200, 40, 40); GREEN = C(40, 150, 85); ORANGE = C(215, 120, 30); WHITE = C(255, 255, 255)

prs = Presentation()
prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
from pptx.util import Inches as I

def slide(): return prs.slides.add_slide(BLANK)

def tx(s, x, y, w, h, text, size=14, bold=False, color=DGRAY, align=PP_ALIGN.LEFT, line=1.2):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = line
        r = p.add_run(); r.text = ln
        f = r.font; f.size = Pt(size); f.bold = bold; f.color.rgb = color; f.name = '微软雅黑'
    return tb

def tit(s, t, sub=""):
    tx(s, 0.5, 0.3, 12.4, 0.7, t, size=26, bold=True, color=BLUE)
    if sub:
        tx(s, 0.5, 0.95, 12.4, 0.45, sub, size=12.5, color=LGRAY)
    bar = s.shapes.add_shape(1, Inches(0.5), Inches(1.0), Inches(12.4), Pt(2.5))
    bar.fill.solid(); bar.fill.fore_color.rgb = BLUE; bar.line.fill.background()

def card(s, x, y, w, h, title, body, col=BLUE, tsize=15, bsize=12.5, fill=None):
    box = s.shapes.add_shape(5, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = fill if isinstance(fill, C) else C(246, 249, 254)
    box.line.color.rgb = col; box.line.width = Pt(1.6)
    tx(s, x + 0.16, y + 0.12, w - 0.32, 0.4, title, size=tsize, bold=True, color=col)
    tx(s, x + 0.16, y + 0.58, w - 0.32, h - 0.7, body, size=bsize, color=DGRAY, line=1.25)

def bullets(s, items, y=1.35, size=13.5, gap=0.52, w=12.3, x=0.62):
    for it in items:
        tx(s, x, y, w, 0.9, it, size=size, color=DGRAY)
        y += gap

def tab(s, x, y, w, rows, fsize=11.5, hrow=0.44):
    from pptx.util import Inches as II
    shp = s.shapes.add_table(len(rows), len(rows[0]), II(x), II(y), II(w), II(hrow * len(rows)))
    tb = shp.table
    for j in range(len(rows[0])):
        tb.columns[j].width = II(w / len(rows[0]))
    for i, row in enumerate(rows):
        tb.rows[i].height = II(hrow)
        for j, v in enumerate(row):
            c = tb.cell(i, j); c.text = str(v)
            for p in c.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
                for r in p.runs:
                    r.font.size = Pt(fsize); r.font.name = '微软雅黑'
                    if i == 0: r.font.bold = True
            if i == 0:
                c.fill.solid(); c.fill.fore_color.rgb = C(217, 226, 243)

# ============ 1 封面
s = slide()
tx(s, 0.9, 1.3, 11.5, 1.3, "磁吸焊接辅助手电", size=44, bold=True, color=BLUE)
tx(s, 0.9, 2.6, 11.5, 1.0, "—— 一个「从零到整机」的开源学习包", size=22, color=DGRAY)
tx(s, 0.9, 3.7, 11.5, 1.4,
   "不是又一块开发板，而是一个完整产品\n硬件 + 固件 + 结构 + 测试，全链路开源 · 可稳定复现 · 低成本",
   size=16, color=LGRAY, line=1.5)
tx(s, 0.9, 5.6, 11.5, 1.0,
   "仿真跑不队：邵子中、李昊桐、姜亚楷　|　指导老师提供 CW32L010 开发板与工程框架\nGitHub：github.com/Davidshao-114514/chuangkexunlianying",
   size=12.5, color=DGRAY, line=1.4)

# ============ 2 为什么做
s = slide(); tit(s, "为什么做这个学习包", "起点：一块老师研发的开发板 与 Arduino 之间的差距")
card(s, 0.55, 1.3, 5.9, 2.3, "开发板 / Arduino 教会我们什么", 
     "• 单个外设怎么用：点灯、按键、刷屏、PWM\n"
     "• 大量现成库与教程，快速看到效果\n"
     "• 但——板载条件已定，不需要做工程决策", col=LGRAY)
card(s, 6.85, 1.3, 5.9, 2.3, "一个完整产品逼着我们回答",
     "• 电池设备为什么敢不用稳压器？那怎么测电压？\n"
     "• 驱动手册写 100Hz–10kHz，到底选多少？\n"
     "• 3 个按键怎么覆盖 8 项功能？\n"
     "• 关机后怎么做到唤醒不花屏、松手不重启？", col=BLUE)
card(s, 0.55, 3.85, 12.2, 1.5, "启发与选择",
     "→ 开发板教你「怎么用芯片」，一个完整产品才教你「怎么做工程决策」。\n"
     "→ 这样的项目不该停在我们三个人的桌面上——它需要开源的共享共建：让别人能稳定复现、并在复现之上继续走远。",
     col=GREEN, tsize=15)
card(s, 0.55, 5.5, 12.2, 1.4, "于是有了这个学习包（三个目标）",
     "① 项目式学习：以「做出一台能用的整机」为目标，覆盖电源/主控/输出/人机/结构/测试六个层面\n"
     "② 稳定复现的低成本结构：通用件 + 3D 打印 + BOM≈¥45 + 14 项量化验收判据\n"
     "③ 共享共建：文档写「为什么」、代码写「怎么」、问题写「坑在哪」",
     col=ORANGE, tsize=15)

# ============ 3 产品是什么
s = slide(); tit(s, "项目是什么：一吸即用、指哪照哪的焊接照明工具",
                 "面向「焊点看得见、双手都解放」的真实场景（磁吸底座 + 万向柔臂 + 双路无级调光）")
card(s, 0.55, 1.35, 3.95, 2.5, "痛点 1 · 遮挡",
     "左手镊子、右手烙铁，光照永远差 5 cm\n\n→ 磁吸 + 万向柔臂：把光贴到焊点上", col=RED)
card(s, 4.7, 1.35, 3.95, 2.5, "痛点 2 · 占位",
     "台灯底座吃掉四分之一桌面\n\n→ 磁吸离场：桌面零占用", col=ORANGE)
card(s, 8.85, 1.35, 3.9, 2.5, "痛点 3 · 不稳/贵",
     "廉价灯夹不稳、专业机床灯 300+ 元\n\n→ 双路独立调光，百元以内", col=GREEN)
tx(s, 0.6, 4.15, 12.2, 2.6,
   "与常见方案的差异（一句话版本）\n"
   "· 相比台灯 / 廉价便携灯：吸上就亮、取下就灭（磁吸 DET 检测，未吸附不输出）\n"
   "· 相比专业机床灯：价格降 75%+，同时保留「万向瞄准 + 无级调光」\n"
   "· 相比「又一块玩具灯」：本项目的价值不止在灯，而在于它**完整可复现的工程体系**——这才是开源学习包的主体",
   size=14, color=DGRAY, line=1.5)

# ============ 4 硬核参数
s = slide(); tit(s, "硬核参数一览（全部有文档依据与踩坑过程）", "对应 docs/03 硬件设计说明、docs/04 固件架构说明")
tab(s, 0.55, 1.35, 12.25, [
    ["项目", "参数", "说明"],
    ["主控", "CW32L010F8（M0+，TSSOP-20）", "HSI 48MHz，无外部晶振，OSC± 复用"],
    ["供电", "锂电池直供 MCU", "无稳压器；CE 控制 3V3 轨（磁吸口）"],
    ["调光", "双路 PWM 8kHz / 6000 级", "ARR=5999，步进 0.5%，下限 3%"],
    ["驱动", "RY3730 ×2", "可 pin-to-pin 替换 SY7200"],
    ["磁吸接口", "4pin：3V3/GND/LED−/DET", "吸附才输出，脱落即灭"],
    ["检测去抖", "100ms（按键仅 15ms）", "去抖时间由物理过程决定"],
    ["按键", "3 键复合语义", "单击/长按（三段加速）/双击"],
    ["电池采样", "内部 BGR 1.2V 反推 VDD", "无外部电路，省 210µA"],
    ["电池保护", "LOW<3350mV / PROTECT<3200mV", "滞回 + 2s 确认 + 告警页"],
    ["显示", "0.5\" CH1115 88×48", "软件 I2C，地址 0x3C"],
    ["掉电记忆", "Flash + RTC 1h 窗口", "记住亮度又不会记住过期设置"],
    ["资源占用", "Flash≈12KB / RAM≈1.1KB", "编译 0 错 0 警告"],
], fsize=11, hrow=0.4)

# ============ 5 系统框图
s = slide(); tit(s, "系统框图：四个域 + 十二类连线", "电源域 / 主控域 / 输出域 / 人机域——任何一部分都可单独替换或升级")
if os.path.exists(BLOCK):
    s.shapes.add_picture(BLOCK, Inches(0.5), Inches(1.25), width=Inches(12.3))
tx(s, 0.55, 6.95, 12.2, 0.5, "完整文件：docs/系统框图.png（可打印）", size=11, color=LGRAY)

# ============ 6 硬件亮点
s = slide(); tit(s, "硬件设计亮点：三个「反常识」决策", "每一个决策都写清了理由与代价（docs/03）")
card(s, 0.55, 1.3, 4.0, 4.6, "① 电池直供 MCU\n（不加稳压器）",
     "为什么：省元件、省静态电流、省板面积\n\n"
     "代价：不能再「分压 + 同源参考」测电压——比例测量会失效\n\n"
     "解法：内部 BGR 1.2V 绝对参考反推 VDD\n"
     "· BGR 仅测量时开启\n"
     "· 成对交替采样（抵消 LED 亮灭漂移）\n"
     "· IIR 平滑（抑制 PWM 纹波）\n\n"
     "省下外部 10K/10K 分压的 210µA 常通耗电", col=BLUE, bsize=12)
card(s, 4.75, 1.3, 4.0, 4.6, "② 磁吸 4pin 与「脱落即灭」",
     "接口：3V3 / GND / LED− / DET\n\n"
     "原理：板侧 DET 下拉（未吸附=低）；灯端把 3V3 接 DET（吸附=高）\n\n"
     "三个收益：\n"
     "· 安全：脱落瞬间切断输出，避免开路/异常驱动\n"
     "· 交互：吸上就亮、取下就灭，无需额外开关\n"
     "· 可诊断：吸附状态上屏，一眼区分「灯坏」与「没吸好」\n\n"
     "去抖 100ms（磁吸撞击远长于人手按键）", col=GREEN, bsize=12)
card(s, 8.95, 1.3, 3.85, 4.6, "③ 频率与亮度下限\n一起设计",
     "调光频率迭代史：\n"
     "20k → 30k → 5k → 1k → 8kHz\n\n"
     "根因：驱动手册推荐 100Hz–10kHz\n\n"
     "最终：8kHz + 亮度下限 3%\n"
     "（低于 3% 会频闪；下限与频率必须一起定）\n\n"
     "→ 分辨率 6000 级，软件用千分比对外暴露，与寄存器解耦", col=ORANGE, bsize=12)

# ============ 7 固件架构
s = slide(); tit(s, "固件架构：分层 BSP + 1ms 分频调度", "12 个 BSP 模块，main.c 只表达产品行为，不知道寄存器细节（docs/04）")
card(s, 0.55, 1.3, 6.1, 4.7, "分层结构",
     "main.c / interrupts_cw32l010.c\n"
     "　应用层：UI 布局 / 按键状态机 / 电池状态机 / 主循环\n"
     "　　　　▲\n"
     "bsp_*.c（12 个模块）\n"
     "　bsp_pwm（8kHz 双路）· bsp_detect（100ms 去抖）\n"
     "　bsp_power（STOP+上升沿唤醒）· bsp_save（Flash 亮度记忆）\n"
     "　bsp_adc（BGR 反推）· bsp_oled（CH1115）· bsp_i2c · bsp_button · bsp_timer\n"
     "　　　　▲\n"
     "CW32L010_Lib（芯片标准库）\n\n"
     "价值：换 MCU 时原则上只重写 bsp_*，main.c 语义不变", col=BLUE, bsize=12.5)
card(s, 6.85, 1.3, 5.95, 2.25, "调度：全部在 1ms 中断内分频",
     "1ms：系统时基 / 告警页关屏倒计时\n"
     "5ms：按键扫描 + 状态机（调光加速/双击判定）+ 吸附检测\n"
     "50ms：ADC 采样 + 过放保护状态机\n"
     "200ms：界面数值刷新请求\n"
     "主循环：只做 OLED 刷屏（耗时 ms 级，不能进中断）", col=GREEN, bsize=12)
card(s, 6.85, 3.75, 5.95, 2.25, "3 键承载 8 项功能",
     "PB03 开关机（含唤醒，上升沿）\n"
     "PA04 亮度−：单击微调 / 长按连续 / 双击切「调节对象」\n"
     "PA05 亮度+：单击微调 / 长按连续 / 双击切「点亮模式」\n"
     "连续调节三段加速：60ms → 1.1s 后 30ms → 1.9s 后 12ms/步\n"
     "有效电平高（外部下拉，抗 PWM 噪声）", col=ORANGE, bsize=12)

# ============ 8 踩坑实录
s = slide(); tit(s, "开发踩坑实录：这些坑比代码更值钱", "docs/08 —— 完整记录「现象 → 根因 → 解决」，可当排错清单用")
tab(s, 0.55, 1.35, 12.25, [
    ["现象", "根因", "解决"],
    ["屏幕全黑（4 轮才点亮）", "CH1115 不认 SSD1306 的 0x8D 命令；方向/偏移不符", "用屏厂官方序列：0x82 + 0xAD 0x8B 0x33 + A0/C0 + D3 0x38"],
    ["电池读数恒为满量程一半", "VBAT 直供 → 分压与参考同源，比例测量失效", "内部 BGR 1.2V 反推 VDD"],
    ["LED 亮起后电压乱跳", "双通道非同时采样撞上 PWM 波动", "成对交替采样 + IIR 平滑"],
    ["低亮度频闪 / 啸叫", "频率脱离驱动手册推荐区（100Hz–10kHz）", "8kHz + 亮度下限 3%"],
    ["关机松手即重启", "按键释放回弹瞬态被当作唤醒沿", "等释放 + 100ms 消抖窗 + 清标志再睡"],
    ["按键误触（PWM 一开就触发）", "高阻节点耦合 PWM 噪声", "硬件下拉电阻改小（软件消抖无效）"],
    ["唤醒后花屏", "屏幕掉电重启需完整初始化", "唤醒延时 500×1ms 后再初始化屏"],
], fsize=10.5, hrow=0.62)
tx(s, 0.6, 6.55, 12.2, 0.7,
   "从这些坑提炼的工程原则：以厂家官方例程为准 · 测电源要有绝对参考 · 频率与占空比一起设计 · 低功耗的坑都在时序上 · 软件治不了的噪声回硬件解决 · 把坑写下来",
   size=12, color=RED, line=1.3)

# ============ 9 低成本复现
s = slide(); tit(s, "稳定复现的低成本结构", "目标：一个没有本项目背景的人，照做能复现出同样的整机（docs/05）")
card(s, 0.55, 1.3, 3.95, 2.6, "成本拆解",
     "PCBA 器件：≈ ¥25–35\n"
     "磁吸机构：≈ ¥8–12\n"
     "外壳/装配（含按键帽）：≈ ¥5–10\n"
     "————————————\n"
     "整机：≈ ¥45 级（不含工具）", col=GREEN, bsize=12.5)
card(s, 4.7, 1.3, 3.95, 2.6, "降本四招",
     "① 通用件优先，不做定制\n"
     "② 立创EDA 工程一键下单打样\n"
     "③ 3D 打印替代开模（含按键帽）\n"
     "④ 驱动 pin-to-pin 可替换（RY3730 ⇄ SY7200）", col=BLUE, bsize=12.5)
card(s, 8.85, 1.3, 3.9, 2.6, "复现路线",
     "采购 → 制板&打印 → 焊接装配 →\n烧录 → 上电自检 → 14 项验收\n\n"
     "配套：完整 BOM（34 行/56 位号）、\n3D 模型、立创EDA 工程、排错表", col=ORANGE, bsize=12.5)
card(s, 0.55, 4.15, 12.2, 2.55, "为什么能「稳定」复现",
     "· 判据可量化：14 项指标（吸附力、漂移、无频闪、联动、静态电流、续航、温升、电量误差、低电保护、亮度记忆…）每项都有测试条件与合格阈值\n"
     "· 记录可复核：验收记录表模板（条件 / 原始数据 / 结论三要素）；失败项必须写排查过程\n"
     "· 坑已排明：12 条真实现象 → 根因 → 解决 的排错表（docs/05 §6）\n"
     "· 版本可对照：固件 / 外壳 / BOM 三者可对上，样机贴版本号；改造项（频率、去抖、窗口、下限）全部做成宏，可配置、可回退",
     col=BLUE, tsize=14, bsize=12.5)

# ============ 10 学习路径
s = slide(); tit(s, "学习路径：9 个模块 / 16–24 学时", "每模块四段：目标 → 读什么 → 动手做什么 → 检查点（docs/02）")
tab(s, 0.55, 1.35, 12.25, [
    ["#", "模块", "你会获得"],
    ["1", "看懂一台整机", "四域分区思维；自己画一张系统框图"],
    ["2", "供电与软开关", "电池直供的取舍；CE 控 3V3；关机为何要 STOP"],
    ["3", "按键与复合语义", "高有效输入、15ms 去抖、单击/长按/双击状态机"],
    ["4", "PWM 与无级调光", "8kHz/6000 级推导；最低占空比下限为何是 3%"],
    ["5", "磁吸检测与安全联动", "复用引脚、100ms 去抖、脱落即灭"],
    ["6", "采样与电池保护状态机", "BGR 绝对参考、成对采样、滞回阈值与告警"],
    ["7", "低功耗与掉电记忆", "唤醒沿/消抖窗/延迟初始化；Flash 记忆时间窗"],
    ["8", "结构、装配与验收", "磁吸底座+按键帽装配；14 项验收判据"],
    ["★", "开发踩坑复盘", "把别人的坑变成自己的经验（最值钱的一节）"],
], fsize=11.5, hrow=0.5)

# ============ 11 共建共享
s = slide(); tit(s, "开源共建共享", "开源不是把代码扔出去，而是让别人能稳定复现、并在此基础上继续走远")
card(s, 0.55, 1.3, 6.0, 2.5, "欢迎的贡献",
     "🐛 报问题：开 Issue（环境 / 现象 / 复现步骤 / 已尝试）\n"
     "🔧 改代码：修 bug、优化去抖、提升可移植性\n"
     "📐 改硬件：替代件验证、布局优化（需附实测数据）\n"
     "🖨 改结构：更好打印的模型、兼容其他磁吸座\n"
     "🎓 教学复用：教师/社团可直接用 docs/07 当课程包\n"
     "🧰 留坑记录：把你的踩坑补进 docs/08（比代码更珍贵）", col=GREEN, bsize=12)
card(s, 6.75, 1.3, 6.05, 2.5, "共建原则（四条）",
     "① 可复现优先：让「下一个人照着做得到同样结果」更容易\n"
     "② 解释为什么：代码与文档都要讲清设计动机\n"
     "③ 小步提交：一次 PR 只做一件事，便于审查与回退\n"
     "④ 友善沟通：本包的目标对象就是初学者", col=BLUE, bsize=12)
card(s, 0.55, 4.05, 12.25, 2.65, "仓库与文档结构",
     "README.md（门面）· CONTRIBUTING.md（共建指南）· LICENSE(MIT)\n"
     "docs/01 为什么做这个学习包　docs/02 学习路径　docs/03 硬件设计说明　docs/04 固件架构说明\n"
     "docs/05 复现指南　docs/06 测试与验收　docs/07 教学设计　docs/08 开发踩坑实录　系统框图.png\n"
     "USER/ 固件（12 个 BSP 模块）· MDK/ Keil 工程 · 工程文件夹/ UserApp完整工程 + 立创EDA + 3D外壳 + BOM\n\n"
     "GitHub：github.com/Davidshao-114514/chuangkexunlianying",
     col=DGRAY, tsize=14, bsize=12)

# ============ 12 结尾
s = slide()
tx(s, 0.9, 2.2, 11.5, 1.4, "一吸即用，指哪照哪", size=40, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
tx(s, 0.9, 3.7, 11.5, 1.2,
   "一个完整产品，抵得过一打外设实验\n欢迎复现、欢迎改进、欢迎把你的坑也留下来",
   size=18, color=DGRAY, align=PP_ALIGN.CENTER, line=1.5)
tx(s, 0.9, 5.3, 11.5, 1.0,
   "仿真跑不队 · 邵子中 / 李昊桐 / 姜亚楷\n感谢指导老师的开发板与工程框架；感谢每一位共建者",
   size=13, color=LGRAY, align=PP_ALIGN.CENTER, line=1.4)

prs.save(FNAME)
print("PPT saved:", FNAME, "| 页数:", len(prs.slides._sldIdLst))
