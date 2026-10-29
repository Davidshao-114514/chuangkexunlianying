# -*- coding: utf-8 -*-
"""全文关键词定位：RACI-V / 五维资源 / 仪表盘 / 资源领域能力"""
import sys, os, glob
from pptx import Presentation

KEYS = ["RACI-V", "五维", "创新核算", "仪表盘", "增值", "自我效能", "动员他人", "财经素养", "资源与调动", "雷达"]
base = "C:/Users/35464/Desktop/创客训练营"
for path in glob.glob(base + "/*.pptx"):
    prs = Presentation(path)
    hits = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for sh in slide.shapes:
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    for c in r.cells:
                        if c.text.strip():
                            texts.append("[T]" + c.text.replace("\n", " / "))
            if sh.has_text_frame and sh.text_frame.text.strip():
                texts.append(sh.text_frame.text)
        body = "\n".join(texts)
        if any(k in body for k in KEYS):
            hits.append((i, body))
    if hits:
        print("#" * 70)
        print("FILE:", os.path.basename(path), "命中页:", [h[0] for h in hits])
        for i, body in hits[:6]:
            print(f"========== Slide {i} ==========")
            print(body[:1800])
