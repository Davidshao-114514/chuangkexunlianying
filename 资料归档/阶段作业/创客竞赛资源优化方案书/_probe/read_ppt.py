# -*- coding: utf-8 -*-
"""提取课件 PPT 全部文本，定位评价量表/资源相关页"""
import sys, os
from pptx import Presentation

for path in sys.argv[1:]:
    print("#" * 70)
    print("FILE:", os.path.basename(path))
    prs = Presentation(path)
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for sh in slide.shapes:
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    for c in r.cells:
                        if c.text.strip():
                            texts.append("[T]" + c.text.replace("\n", " / "))
            if sh.has_text_frame and sh.text_frame.text.strip():
                texts.append(sh.text_frame.text)
        body = "\n".join(texts)
        if any(k in body for k in ("量表", "权重", "资源", "雷达", "RACI", "评价", "财经")):
            print(f"========== Slide {i} ==========")
            print(body[:2500])
