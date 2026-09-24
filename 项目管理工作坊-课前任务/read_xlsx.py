# -*- coding: utf-8 -*-
import openpyxl
p = r'C:\Users\35464\Desktop\创客训练营\痛点知识领域映射表.xlsx'
wb = openpyxl.load_workbook(p, data_only=True)
for name in wb.sheetnames:
    ws = wb[name]
    print(f'===== Sheet: {name} ({ws.max_row}行 x {ws.max_column}列) =====')
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 60), values_only=True):
        vals = [str(v) if v is not None else '' for v in row]
        if any(vals):
            print(' | '.join(vals))
