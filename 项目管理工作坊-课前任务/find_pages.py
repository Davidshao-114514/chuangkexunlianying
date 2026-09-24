# -*- coding: utf-8 -*-
from pptx import Presentation
p = r'C:\Users\35464\AppData\Local\hermes\attachments\创客训练营——项目管理工作坊——2026创客1班.pptx'
prs = Presentation(p)
for i, s in enumerate(prs.slides, 1):
    if not (82 <= i <= 98):
        continue
    print(f'===== P{i} =====')
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            t = sh.text_frame.text.strip().replace('\n', ' ⏎ ')
            print('  [T]', t[:1000])
        if sh.has_table:
            for row in sh.table.rows:
                print('  [R]', ' | '.join(c.text.replace('\n', ' ')[:70] for c in row.cells)[:500])
