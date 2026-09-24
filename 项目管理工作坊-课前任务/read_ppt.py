# -*- coding: utf-8 -*-
from pptx import Presentation
p = r'C:\Users\35464\AppData\Local\hermes\attachments\创客训练营——项目管理工作坊——2026创客1班.pptx'
prs = Presentation(p)
print('幻灯片数:', len(prs.slides))
for i, s in enumerate(prs.slides, 1):
    txts = []
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            txts.append(sh.text_frame.text.strip().replace('\n', ' ⏎ '))
        if sh.has_table:
            for row in sh.table.rows:
                txts.append('[表] ' + ' | '.join(c.text.replace('\n',' ')[:50] for c in row.cells))
    if txts:
        print(f'===== P{i} =====')
        for t in txts[:12]:
            print(' ', t[:500])
